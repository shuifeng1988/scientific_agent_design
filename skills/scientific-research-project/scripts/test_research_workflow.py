"""Small isolated acceptance fixtures. All data/search/approval records are synthetic.

No GPU, network search, model inference, real approval or scientific claim is made.
"""
import copy
import csv
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

import research_workflow as rw
import project_events as pe
import supervisor_daemon as sd


def build_project(root, frozen=True):
    root = Path(root)
    rw.draft(root, dict(project="# Draft fixture question Q01\n", plan="# Draft steps for fixture\n"))
    for name in ("search_a", "search_b", "user_message"):
        rw.atomic(root / ("provenance/" + name + ".txt"), "SYNTHETIC TEST FIXTURE ONLY: " + name)
    searches = [dict(route=name, query="synthetic fixture query", date="2026-09-09",
                    recent_window="24 months", selection_rules="fixture only, no scientific evidence",
                    evidence=rw.descriptor(root / ("provenance/" + name + ".txt"), root))
                for name in ("search_a", "search_b")]
    appraisal = {k: "SYNTHETIC FIXTURE: " + k for k in ("question", "searched_at", "significance",
                    "novelty", "feasibility", "worth_doing", "information_gain", "limitations", "recommendation")}
    appraisal.update(searches=searches, sources=[dict(url="https://example.org/fixture",
                     title="Synthetic fixture", year="2026", evidence_direction="fixture", relevance="test only")],
                     methods={k: "fixture " + k for k in ("conventional", "current_reproducible", "frontier")})
    rw.atomic(root / "provenance/appraisal.json", appraisal)
    tasks = [dict(task_id="Q01.0" + str(i), question_id="Q01", title="Fixture task " + str(i),
                  result_path="results/Q01/Q01.0" + str(i), required_models="fixture", required_data="fixture")
             for i in (1, 2)]
    docs = {name: "# SYNTHETIC workflow test\n" for name in rw.ROOT_DOCS}
    docs["00_PROJECT.md"] += "Q01: Test plan-to-results correspondence. No scientific claim.\n"
    docs["01_PLAN.md"] += "\n".join(t["task_id"] + " " + t["result_path"] for t in tasks)
    proposal = dict(version="fixture-v1", documents=docs, tasks=tasks,
                    appraisal=rw.descriptor(root / "provenance/appraisal.json", root))
    proposal_path = root / "provenance/proposal.json"
    rw.atomic(proposal_path, proposal)
    approval_path = root / "provenance/approval.json"
    rw.atomic(approval_path, dict(decision="approved", proposal_sha256=rw.sha(proposal_path),
              appraisal_sha256=proposal["appraisal"]["sha256"], actor="TEST_FIXTURE_NOT_HUMAN",
              message="Synthetic approval used only for tests", recorded_at="2026-09-09",
              evidence=rw.descriptor(root / "provenance/user_message.txt", root)))
    if frozen:
        rw.freeze(root, proposal_path, approval_path)
    return proposal_path, approval_path


def prepare_report(root, task_id, run_id, language="en"):
    from PIL import Image, ImageDraw
    root = Path(root); result = root / rw.task(root, task_id)["result_path"]
    result.mkdir(parents=True, exist_ok=True)
    # Real small computation on explicitly synthetic data.
    values = [1, 2, 3, 4]
    count, mean = len(values), sum(values) / len(values)
    rw.atomic(root / "data/fixture.csv", "value\n" + "\n".join(map(str, values)) + "\n")
    rw.atomic(result / "tables/summary.csv", f"metric,value\ncount,{count}\nmean,{mean}\n")
    (result / "figures").mkdir(exist_ok=True)
    image = Image.new("RGB", (500, 300), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 15), f"SYNTHETIC FIXTURE: mean={mean}, count={count}", fill="black")
    draw.line((50, 250, 450, 250), fill="black")
    draw.line((50, 250, 50, 50), fill="black")
    draw.rectangle((170, 250 - 50 * mean, 270, 250), fill="steelblue")
    draw.text((190, 255), "mean", fill="black")
    draw.text((15, 250 - 50 * mean - 5), str(mean), fill="black")
    image.save(result / "figures/mean.png")
    sections = {k: "Fixture " + k + ": no scientific inference." for k in rw.SECTIONS}
    sections["results"] = f"Synthetic sample count={count}; mean={mean}."
    sections["conclusions"] = f"The fixture arithmetic mean is {mean}; no real research conclusion."
    if language == "zh":
        sections = {k: "合成测试：" + k + "，不代表科研结论。" for k in rw.SECTIONS}
        sections["results"] = f"合成样本数为 {count}，均值为 {mean}。"
    manifest = dict(task_id=task_id, run_id=run_id, plan_version="fixture-v1", language=language,
          title="SYNTHETIC WORKFLOW ACCEPTANCE", host="local", environment="Python unit-test fixture",
          command=[sys.executable, "-B", str(Path(__file__).resolve()), "fixture", str(root), task_id, "run"],
          software=[dict(name="Python", version=sys.version.split()[0],
          usage="Compute mean of four synthetic numbers")], sections=sections,
          inputs=[dict(**rw.descriptor(root / "data/fixture.csv", root), source="synthetic test generator",
                       version="fixture-v1", access="project-local fixture")],
          tables=[dict(**rw.descriptor(result / "tables/summary.csv", result), caption="Synthetic summary")],
          figures=[dict(**rw.descriptor(result / "figures/mean.png", result), caption="Synthetic mean",
                        source_tables=["tables/summary.csv"], method="Pillow fixture bar chart")])
    mp = result / "_evidence/report.json"
    rw.atomic(mp, manifest)
    return mp


def finish_report(root, task_id, run_id):
    root = Path(root); result = root / rw.task(root, task_id)["result_path"]
    reflection = {k: "Synthetic fixture only" for k in ("expected", "observed", "uncertainty", "reasoning",
                    "alternative_explanations", "evidence_strength", "limitations", "next_step_impact")}
    reflection.update(run_id=run_id, decision="continue", literature_search=dict(
        queries=["synthetic fixture"], searched_at="2026-09-09", status="no_relevant_results",
        fixture=True, sources=[]))
    rw.atomic(result / "_evidence/reflection.json", reflection)
    r = read_receipt = rw.read(result / "_evidence/report_receipt.json")
    paths = [a["path"] for a in r["artifacts"]] + [r["manifest"]["path"],
             "_evidence/report_receipt.json", "_evidence/reflection.json"]
    rw.atomic(result / "_evidence/qc.json", dict(run_id=run_id, passed=True, fixture=True,
              artifacts=[rw.descriptor(result / p, result) for p in paths]))
    return result


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
    def tearDown(self):
        self.temp.cleanup()

    def test_draft_only_two_roots(self):
        build_project(self.root, frozen=False)
        self.assertTrue((self.root / "00_PROJECT.md").exists())
        self.assertFalse((self.root / "02_SOFTWARE.md").exists())

    def test_missing_or_stale_user_approval_blocks_freeze(self):
        p, a = build_project(self.root, frozen=False)
        approval = rw.read(a); approval["decision"] = "pending"; rw.atomic(a, approval)
        with self.assertRaisesRegex(ValueError, "user approval"):
            rw.freeze(self.root, p, a)
        approval["decision"] = "approved"; approval["proposal_sha256"] = "0" * 64; rw.atomic(a, approval)
        with self.assertRaisesRegex(ValueError, "cover proposal"):
            rw.freeze(self.root, p, a)
        self.assertFalse((self.root / "registry/design.json").exists())

    def test_missing_search_evidence_blocks_freeze(self):
        p, a = build_project(self.root, frozen=False)
        (self.root / "provenance/search_b.txt").unlink()
        with self.assertRaisesRegex(ValueError, "hash mismatch"):
            rw.freeze(self.root, p, a)

    def test_freeze_five_documents_and_provenance(self):
        build_project(self.root)
        self.assertEqual(rw.verify_design(self.root)["tasks"], 2)
        for n in rw.ROOT_DOCS:
            self.assertTrue((self.root / n).is_file())

    def test_preliminary_appraisal_cannot_be_newly_frozen(self):
        p, a = build_project(self.root, frozen=False)
        appraisal_path = self.root / "provenance/appraisal.json"
        appraisal = rw.read(appraisal_path)
        appraisal["status"] = "preliminary_not_approved"
        rw.atomic(appraisal_path, appraisal)
        proposal = rw.read(p)
        proposal["appraisal"] = rw.descriptor(appraisal_path, self.root)
        rw.atomic(p, proposal)
        approval = rw.read(a)
        approval.update(proposal_sha256=rw.sha(p), appraisal_sha256=rw.sha(appraisal_path))
        rw.atomic(a, approval)
        with self.assertRaisesRegex(ValueError, "not ready for design freeze"):
            rw.freeze(self.root, p, a)
        self.assertFalse((self.root / "registry/design.json").exists())

    def test_ready_appraisal_still_requires_exact_user_approval(self):
        p, a = build_project(self.root, frozen=False)
        appraisal_path = self.root / "provenance/appraisal.json"
        appraisal = rw.read(appraisal_path)
        appraisal["status"] = "ready_for_design_review"
        rw.atomic(appraisal_path, appraisal)
        proposal = rw.read(p)
        proposal["appraisal"] = rw.descriptor(appraisal_path, self.root)
        rw.atomic(p, proposal)
        with self.assertRaisesRegex(ValueError, "cover proposal"):
            rw.freeze(self.root, p, a)
        approval = rw.read(a)
        approval.update(proposal_sha256=rw.sha(p), appraisal_sha256=rw.sha(appraisal_path))
        rw.atomic(a, approval)
        self.assertTrue(rw.freeze(self.root, p, a)["passed"])
        self.assertTrue((self.root / "scripts/hooks/project_event.py").is_file())

    def test_plan_change_rejected(self):
        build_project(self.root)
        with (self.root / "01_PLAN.md").open("a") as f: f.write("changed endpoint")
        with self.assertRaisesRegex(ValueError, "frozen scientific"):
            rw.verify_design(self.root)

    def test_registry_path_change_rejected(self):
        build_project(self.root)
        rows = rw.registry(self.root); rows[0]["result_path"] += "_wrong"
        rw.atomic(self.root / "registry/tasks.tsv", rw.tasks_table(rows))
        with self.assertRaisesRegex(ValueError, "mapping changed"):
            rw.verify_design(self.root)

    def test_input_tamper_blocks_report(self):
        build_project(self.root)
        mp = prepare_report(self.root, "Q01.01", "Q01.01.R001")
        rw.atomic(self.root / "data/fixture.csv", "tampered")
        with self.assertRaisesRegex(ValueError, "hash mismatch"):
            rw.render_report(self.root, "Q01.01", mp)

    def test_report_generation_and_bilingual_content(self):
        build_project(self.root)
        for tid, lang in (("Q01.01", "en"), ("Q01.02", "zh")):
            run = tid + ".R001"
            mp = prepare_report(self.root, tid, run, lang)
            rw.render_report(self.root, tid, mp)
            finish_report(self.root, tid, run)
            self.assertTrue(rw.validate_package(self.root, tid, rw.task(self.root, tid)["result_path"], run)["passed"])

    def test_figure_requires_source_table(self):
        build_project(self.root)
        mp = prepare_report(self.root, "Q01.01", "Q01.01.R001")
        m = rw.read(mp); m["figures"][0]["source_tables"] = []; rw.atomic(mp, m)
        with self.assertRaisesRegex(ValueError, "lineage"):
            rw.render_report(self.root, "Q01.01", mp)

    def package(self):
        build_project(self.root); run = "Q01.01.R001"
        mp = prepare_report(self.root, "Q01.01", run)
        rw.render_report(self.root, "Q01.01", mp)
        return finish_report(self.root, "Q01.01", run)

    def test_zero_byte_word_cannot_pass(self):
        result = self.package()
        (result / "REPORT.docx").write_bytes(b"")
        with self.assertRaises(ValueError):
            rw.validate_package(self.root, "Q01.01", "results/Q01/Q01.01", "Q01.01.R001")

    def test_fake_markdown_value_fails_even_with_rehashed_files(self):
        result = self.package()
        p = result / 'README.md'
        rw.atomic(p, p.read_text().replace('| mean | 2.5 |', '| mean | 99 |'))
        r = rw.read(result / '_evidence/report_receipt.json')
        for a in r['artifacts']:
            if a['path'] == 'README.md':
                a['sha256'] = rw.sha(p)
        rw.atomic(result / '_evidence/report_receipt.json', r)
        finish_report(self.root, 'Q01.01', 'Q01.01.R001')
        with self.assertRaisesRegex(ValueError, 'Markdown/source table'):
            rw.validate_package(self.root, 'Q01.01', 'results/Q01/Q01.01', 'Q01.01.R001')

    def test_fake_word_value_fails_even_with_rehashed_files(self):
        from docx import Document
        result = self.package()
        p = result / 'REPORT.docx'
        d = Document(p); d.tables[0].rows[2].cells[1].text = '99'; d.save(p)
        r = rw.read(result / '_evidence/report_receipt.json')
        for a in r['artifacts']:
            if a['path'] == 'REPORT.docx':
                a['sha256'] = rw.sha(p)
        rw.atomic(result / '_evidence/report_receipt.json', r)
        finish_report(self.root, 'Q01.01', 'Q01.01.R001')
        with self.assertRaisesRegex(ValueError, 'Word/source table'):
            rw.validate_package(self.root, 'Q01.01', 'results/Q01/Q01.01', 'Q01.01.R001')

    def test_existing_report_not_silently_overwritten(self):
        result = self.package()
        with self.assertRaisesRegex(ValueError, 'already rendered'):
            rw.render_report(self.root, 'Q01.01', result / '_evidence/report.json')

    def test_software_event_updates_document(self):
        build_project(self.root)
        rw.atomic(self.root / 'software/version.txt', 'fixture software 1.0')
        pe.emit(self.root, dict(event_type='software.smoke.passed', entity_id='fixture',
                verified=True, message='fixture version check',
                **rw.descriptor(self.root / 'software/version.txt', self.root), meta={}))
        self.assertIn('software/version.txt', (self.root / '02_SOFTWARE.md').read_text())

    def test_report_cells_with_pipe_and_newline(self):
        build_project(self.root)
        mp = prepare_report(self.root, 'Q01.01', 'Q01.01.R001')
        result = mp.parent.parent
        rw.atomic(result / 'tables/summary.csv', 'metric,value\n"a|b","two\nlines"\n')
        m = rw.read(mp)
        m['tables'][0]['sha256'] = rw.sha(result / 'tables/summary.csv')
        rw.atomic(mp, m)
        rw.render_report(self.root, 'Q01.01', mp)
        finish_report(self.root, 'Q01.01', 'Q01.01.R001')
        self.assertTrue(rw.validate_package(self.root, 'Q01.01', 'results/Q01/Q01.01', 'Q01.01.R001')['passed'])

    def test_missing_design_never_falls_back_to_legacy(self):
        build_project(self.root)
        rw.atomic(self.root / 'registry/supervisor.json', fixture_config(self.root))
        (self.root / 'registry/design.json').unlink()
        def driver(host, action, payload):
            self.assertEqual(action, 'resources')
            return dict(timestamp=time.time(), free_cpus=4, free_memory_gb=8, gpus=[])
        with patch.object(sd, 'driver', side_effect=driver), patch.object(sd.Supervisor, 'sync_monitor'):
            s = sd.Supervisor(self.root); s.tick()
        self.assertIn('workflow_hold', s.state)
        self.assertTrue(all(n['attempt'] == 0 for n in s.state['nodes'].values()))

    def test_partial_qc_cannot_pass(self):
        result = self.package()
        q = rw.read(result / "_evidence/qc.json"); q["artifacts"] = q["artifacts"][:1]
        rw.atomic(result / "_evidence/qc.json", q)
        with self.assertRaisesRegex(ValueError, "complete report evidence"):
            rw.validate_package(self.root, "Q01.01", "results/Q01/Q01.01", "Q01.01.R001")

    def test_wrong_run_or_task_rejected(self):
        self.package()
        with self.assertRaisesRegex(ValueError, "run/task"):
            rw.validate_package(self.root, "Q01.01", "results/Q01/Q01.01", "OLD.R001")
        with self.assertRaisesRegex(ValueError, "canonical"):
            rw.validate_package(self.root, "Q01.02", "results/Q01/Q01.01", "Q01.01.R001")

    def test_event_idempotency_and_root_doc_update(self):
        build_project(self.root)
        rw.atomic(self.root / "data/download.txt", "fixture download")
        event = dict(event_type="data.download.completed", entity_id="fixture",
                     verified=True, message="fixture download verified",
                     **rw.descriptor(self.root / "data/download.txt", self.root), meta={})
        pe.emit(self.root, event); out = pe.emit(self.root, event)
        self.assertTrue(out["duplicate"])
        events = [json.loads(line) for line in (self.root / 'provenance/events.jsonl').read_text().splitlines()]
        self.assertEqual(sum(e['event_type'] == 'data.download.completed' for e in events), 1)
        self.assertIn("data/download.txt", (self.root / "03_DATA.md").read_text())
        self.assertTrue(rw.verify_design(self.root)["passed"])

    def test_no_completion_before_qc(self):
        build_project(self.root)
        with self.assertRaises(ValueError):
            pe.emit(self.root, dict(event_type="run.progress", task_id="Q01.01", status="complete",
                    message="false completion", meta={}))
        self.assertEqual(rw.task(self.root, "Q01.01")["status"], "pending")

    def test_full_publication_and_duplicate_hook(self):
        self.package()
        event = dict(event_type="qc.passed", task_id="Q01.01", status="complete", verified=True,
                     path="results/Q01/Q01.01", message="fixture only", meta={})
        pe.emit(self.root, event); pe.emit(self.root, event)
        self.assertEqual(rw.task(self.root, "Q01.01")["status"], "complete")
        self.assertIn("Q01.01/README.md", (self.root / "results/README.md").read_text())

    def test_scientific_anomaly_blocks_publication(self):
        result = self.package()
        rev = rw.read(result / "_evidence/reflection.json"); rev["decision"] = "needs_user"
        rw.atomic(result / "_evidence/reflection.json", rev)
        q = rw.read(result / "_evidence/qc.json")
        for x in q["artifacts"]:
            if x["path"] == "_evidence/reflection.json":
                x["sha256"] = rw.sha(result / x["path"])
        rw.atomic(result / "_evidence/qc.json", q)
        with self.assertRaisesRegex(ValueError, "scientific hold"):
            pe.emit(self.root, dict(event_type="qc.passed", task_id="Q01.01", status="complete",
                verified=True, path="results/Q01/Q01.01", message="fixture", meta={}))

    def test_daemon_reconciles_but_holds_new_launch_when_design_changes(self):
        build_project(self.root)
        config = fixture_config(self.root)
        rw.atomic(self.root / "registry/supervisor.json", config)
        jobs, launches = {}, []
        def driver(host, action, payload):
            if action == "resources":
                return dict(timestamp=time.time(), free_cpus=4, free_memory_gb=8, gpus=[])
            if action == "launch":
                launches.append(payload["job_id"]); jobs[payload["job_id"]] = dict(state="running")
                return dict(job_id=payload["job_id"])
            return jobs[payload["job_id"]]
        with patch.object(sd, "driver", side_effect=driver), patch.object(sd.Supervisor, "sync_monitor"):
            s = sd.Supervisor(self.root); s.tick()
            self.assertEqual(len(launches), 1)
            rw.atomic(self.root / "01_PLAN.md", "changed plan")
            s.tick()
            self.assertEqual(s.state["nodes"]["Q01.01"]["status"], "running")
            self.assertEqual(len(launches), 1)
            self.assertIn("workflow_hold", s.state)

    def test_word_zoom_has_schema_required_percent(self):
        from docx import Document
        from docx.oxml.ns import qn
        result = self.package()
        zoom = Document(result / 'REPORT.docx').settings.element.find(qn('w:zoom'))
        self.assertEqual(zoom.get(qn('w:percent')), '100')

    def test_qc_failure_creates_amendment_draft_without_changing_plan(self):
        build_project(self.root)
        before = rw.sha(self.root / "01_PLAN.md")
        event = dict(event_type="qc.failed", task_id="Q01.01", message="fixture contradiction", meta={})
        pe.emit(self.root, event); pe.emit(self.root, event)
        drafts = list((self.root / "provenance/amendment_drafts").glob("*.json"))
        self.assertEqual(len(drafts), 1)
        self.assertEqual(rw.read(drafts[0])["state"], "draft")
        self.assertEqual(rw.task(self.root, "Q01.01")["status"], "blocked")
        self.assertEqual(rw.sha(self.root / "01_PLAN.md"), before)

    def test_cli_check_enforces_frozen_design(self):
        build_project(self.root)
        rw.atomic(self.root / "registry/supervisor.json", fixture_config(self.root))
        argv = [sys.executable, "-B", str(Path(sd.__file__)), "check", "--root", str(self.root)]
        good = subprocess.run(argv, capture_output=True, text=True, timeout=10)
        self.assertEqual(good.returncode, 0, good.stderr)
        self.assertTrue(json.loads(good.stdout)["workflow"]["passed"])
        rw.atomic(self.root / "01_PLAN.md", "unapproved change")
        bad = subprocess.run(argv, capture_output=True, text=True, timeout=10)
        self.assertNotEqual(bad.returncode, 0)

    def test_bootstrap_requires_reviewed_hosts_and_keeps_strict_policy(self):
        build_project(self.root)
        rw.atomic(self.root / "registry/dependencies.tsv",
                  "task_id\tdepends_on\tfrozen_evidence\nQ01.01\t\t\nQ01.02\tQ01.01\t\n")
        hosts = fixture_config(self.root)["hosts"]
        rw.atomic(self.root / "registry/hosts.json", hosts)
        argv = [sys.executable, "-B", str(Path(__file__).with_name("bootstrap_supervisor.py")),
                "--root", str(self.root), "--dependencies", str(self.root / "registry/dependencies.tsv")]
        bad = subprocess.run(argv, capture_output=True, text=True, timeout=10)
        self.assertNotEqual(bad.returncode, 0)
        good = subprocess.run(argv + ["--hosts", str(self.root / "registry/hosts.json")],
                              capture_output=True, text=True, timeout=10)
        self.assertEqual(good.returncode, 0, good.stderr)
        config = rw.read(self.root / "registry/supervisor.json")
        self.assertEqual(config["hosts"], hosts)
        self.assertEqual(config["workflow_policy"], "plan_results_v1")
        self.assertEqual(config["nodes"][1]["depends_on"], ["Q01.01"])
        self.assertTrue(all(n["contract"] is None for n in config["nodes"]))

    def test_readme_structural_parity(self):
        import re
        repo = Path(__file__).resolve().parents[3]
        if not (repo / "README.zh-CN.md").exists():
            self.skipTest("Source README check; installed Skill does not include source-root docs")
        zh, en = [(repo / f"README.{lang}.md").read_text() for lang in ("zh-CN", "en")]
        headings = lambda s: re.findall(r"^#{2,3} (\d+(?:\.\d+)*)", s, re.M)
        self.assertEqual(headings(zh), headings(en))
        self.assertEqual(re.findall(r"```bash\n(.*?)```", zh, re.S),
                         re.findall(r"```bash\n(.*?)```", en, re.S))
        tables = lambda s: [len(t.splitlines()) for t in re.findall(r"(?:^\|.*\n)+", s, re.M)]
        self.assertEqual(tables(zh), tables(en))
        graphs = lambda s: re.sub(r'"[^"]*"', '"LABEL"', re.search(r"```mermaid\n(.*?)```", s, re.S)[1])
        self.assertEqual(graphs(zh), graphs(en))
        self.assertIn((repo / "VERSION").read_text().strip(), zh)
        self.assertIn((repo / "VERSION").read_text().strip(), en)

    def test_real_subprocess_report_chain_and_resume(self):
        build_project(self.root)
        rw.atomic(self.root / "registry/supervisor.json", fixture_config(self.root))
        jobs, launches = {}, []
        def driver(host, action, p):
            if action == "resources":
                return dict(timestamp=time.time(), free_cpus=4, free_memory_gb=8, gpus=[])
            if action == "launch":
                if p["job_id"] in jobs:
                    return dict(job_id=p["job_id"])
                launches.append(p["job_id"])
                env = dict(os.environ, SCI_RUN_ID=p["run_id"], PYTHONDONTWRITEBYTECODE="1")
                done = subprocess.run(p["spec"]["argv"], env=env, capture_output=True, text=True, timeout=20)
                if done.returncode:
                    raise RuntimeError(done.stderr)
                jobs[p["job_id"]] = dict(state="succeeded")
                return dict(job_id=p["job_id"])
            return jobs.get(p["job_id"], dict(state="not_found"))
        with patch.object(sd, "driver", side_effect=driver), patch.object(sd.Supervisor, "sync_monitor"):
            for _ in range(10):
                s = sd.Supervisor(self.root); s.tick()  # durable restart each step
                if all(n["status"] == "complete" for n in s.state["nodes"].values()):
                    break
        self.assertEqual(len(launches), 6, json.dumps(s.state, indent=2))  # no duplicate runs
        self.assertTrue(all(n["status"] == "complete" for n in s.state["nodes"].values()))
        self.assertEqual([r["status"] for r in rw.registry(self.root)], ["complete", "complete"])
        self.assertTrue((self.root / "results/Q01/Q01.02/REPORT.docx").exists())


def fixture_config(root):
    script = Path(__file__).resolve()
    nodes = []
    for i in (1, 2):
        tid = "Q01.0" + str(i)
        contract = dict(gates=[dict(path="00_PROJECT.md", sha256=rw.sha(root / "00_PROJECT.md"))])
        for phase in sd.PHASES:
            contract[phase] = dict(host="local", lightweight=True, cwd=str(root),
                  argv=[sys.executable, "-B", str(script), "fixture", str(root), tid, phase],
                  timeout_seconds=30, resources=dict(cpus=0.2, memory_gb=0.25, gpus=0, gpu_memory_gb=0))
        nodes.append(dict(id=tid, task_id=tid, result_path=rw.task(root, tid)["result_path"],
                   depends_on=[] if i == 1 else ["Q01.01"], scientific_result=True,
                   finalizes_task=True, contract=contract, approved_contract_sha256=sd.digest(contract)))
    return dict(execution_policy="v2", workflow_policy="plan_results_v1", hosts={"local": dict(driver=["fixture"],
                       limits=dict(max_jobs=2, cpus=2, memory_gb=4))}, nodes=nodes)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "fixture":
        root, tid, phase = Path(sys.argv[2]), sys.argv[3], sys.argv[4]
        run = os.environ["SCI_RUN_ID"]
        if phase == "run":
            mp = prepare_report(root, tid, run)
            rw.render_report(root, tid, mp)
        elif phase == "validate":
            finish_report(root, tid, run)
        else:
            rw.validate_package(root, tid, rw.task(root, tid)["result_path"], run)
    else:
        unittest.main()
