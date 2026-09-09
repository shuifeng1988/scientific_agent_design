"""Plan-to-results workflow gates and deterministic Markdown/Word reporting.

Scientific content and user approval must be supplied by real workers/operators.
This module validates evidence and correspondence; it does not establish truth.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import re
import tempfile
import zipfile

ROOT_DOCS = ("00_PROJECT.md", "01_PLAN.md", "02_SOFTWARE.md", "03_DATA.md", "04_STATUS.md")
TASK_FIELDS = ("task_id", "question_id", "title", "result_path", "required_models", "required_data")
REGISTRY_FIELDS = TASK_FIELDS + ("status", "next_action")
SECTIONS = ("data", "methods", "results", "conclusions", "interpretation", "limitations", "qc", "next_step")
HEADINGS = {
    "en": ("Data sources", "Methods", "Results", "Conclusions", "Interpretation",
           "Limitations", "Quality control", "Next step"),
    "zh": ("数据来源", "方法", "结果", "结论", "解释", "局限", "质量控制", "下一步"),
}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def text(value, label):
    require(isinstance(value, str) and bool(value.strip()), "missing " + label)
    return value


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for b in iter(lambda: f.read(1048576), b""):
            h.update(b)
    return h.hexdigest()


def atomic(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    value = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2) + "\n"
    fd, temp = tempfile.mkstemp(prefix=".workflow-", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(value)
            f.flush()
            os.fsync(f.fileno())
        os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def table_line(row):
    return '| ' + ' | '.join(v.replace('|', chr(92) + '|').replace('\n', ' ') for v in row) + ' |'


def inside(root, value):
    root = Path(root).resolve()
    require(isinstance(value, str) and not Path(value).is_absolute(), "project-relative path required")
    p = (root / value).resolve()
    require(p != root and p.is_relative_to(root), "path escapes root: " + value)
    return p


def receipt(root, value):
    p = inside(root, value["path"])
    require(p.is_file() and sha(p) == value["sha256"], "evidence hash mismatch: " + value["path"])
    return p


def descriptor(path, root):
    return dict(path=str(Path(path).resolve().relative_to(Path(root).resolve())), sha256=sha(path))


def tasks_table(tasks):
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=REGISTRY_FIELDS, delimiter="\t", lineterminator="\n")
    writer.writeheader()
    for task in tasks:
        writer.writerow({k: task.get(k, "pending" if k == "status" else "") for k in REGISTRY_FIELDS})
    return out.getvalue()


def registry(root):
    with (Path(root) / "registry/tasks.tsv").open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f, delimiter="\t"))
    require(len({r["task_id"] for r in rows}) == len(rows), "duplicate registered tasks")
    return rows


def validate_tasks(root, tasks, project, plan):
    require(isinstance(tasks, list) and tasks, "nonempty task list required")
    ids, paths = set(), set()
    for t in tasks:
        for k in TASK_FIELDS:
            require(isinstance(t.get(k), str), "task field required: " + k)
            require('\t' not in t[k] and '\n' not in t[k], 'task fields must be single-line')
        require(re.fullmatch(r"Q[0-9]+\.[0-9]+", t["task_id"]), "invalid task ID")
        require(t["task_id"].rsplit(".", 1)[0] == t["question_id"], "question/task mismatch")
        text(t["title"], "task title")
        require(t["question_id"] in project, "question absent from project")
        require(t["task_id"] in plan and t["result_path"] in plan, "task/path absent from plan")
        p = inside(root, t["result_path"])
        require(p.is_relative_to((Path(root) / "results").resolve()), "task path outside results")
        require(t["task_id"] not in ids and p not in paths, "duplicate task or path")
        require(not any(p.is_relative_to(q) or q.is_relative_to(p) for q in paths),
                "canonical task paths cannot be nested")
        ids.add(t["task_id"])
        paths.add(p)


def validate_appraisal(root, ref):
    a = read(receipt(root, ref))
    for k in ("question", "searched_at", "significance", "novelty", "feasibility",
              "worth_doing", "information_gain", "limitations", "recommendation"):
        text(a.get(k), "appraisal." + k)
    require(isinstance(a.get("searches"), list) and len(a["searches"]) >= 2, "two search routes required")
    routes = set()
    for s in a["searches"]:
        for k in ("route", "query", "date", "recent_window", "selection_rules"):
            text(s.get(k), "search." + k)
        receipt(root, s["evidence"])
        routes.add(s["route"])
    require(len(routes) >= 2, "distinct search routes required")
    require(a.get("sources"), "primary source records required")
    for s in a["sources"]:
        for k in ("url", "title", "year", "evidence_direction", "relevance"):
            text(s.get(k), "source." + k)
        require(s["url"].startswith(("https://", "http://")), "source URL required")
    for tier in ("conventional", "current_reproducible", "frontier"):
        text(a.get("methods", {}).get(tier), "method tier " + tier)
    return a


def draft(root, spec):
    root = Path(root).resolve()
    require(not any((root / n).exists() for n in ROOT_DOCS), "draft would overwrite existing project")
    project = text(spec.get("project"), "project draft")
    plan = text(spec.get("plan"), "plan draft")
    root.mkdir(parents=True, exist_ok=True)
    atomic(root / "00_PROJECT.md", project)
    atomic(root / "01_PLAN.md", plan)
    atomic(root / "provenance/intake.json", dict(state="draft", scientific_result=False,
                                               project_sha256=sha(root / ROOT_DOCS[0]),
                                               plan_sha256=sha(root / ROOT_DOCS[1])))
    return dict(state="draft", documents=list(ROOT_DOCS[:2]))


def freeze(root, proposal_path, approval_path):
    root = Path(root).resolve()
    require((root / "provenance/intake.json").is_file(), "intake draft required before freeze")
    require(not (root / "registry/design.json").exists(), "design already frozen; explicit amendment required")
    require(not (root / "registry/tasks.tsv").exists(), "existing registry requires reviewed migration")
    require(not (root / "scripts/hooks/project_event.py").exists(), "existing hook requires reviewed migration")
    proposal_path = Path(proposal_path).resolve()
    approval_path = Path(approval_path).resolve()
    require(proposal_path.is_relative_to(root) and approval_path.is_relative_to(root),
            "proposal and approval must live in project provenance")
    p, approval = read(proposal_path), read(approval_path)
    text(p.get("version"), "plan version")
    require(set(p.get("documents", {})) == set(ROOT_DOCS), "exactly five final documents required")
    for name, value in p["documents"].items():
        text(value, name)
    validate_tasks(root, p["tasks"], p["documents"][ROOT_DOCS[0]], p["documents"][ROOT_DOCS[1]])
    validate_appraisal(root, p["appraisal"])
    require(approval.get("decision") == "approved", "explicit user approval required")
    require(approval.get("proposal_sha256") == sha(proposal_path), "approval does not cover proposal")
    require(approval.get("appraisal_sha256") == p["appraisal"]["sha256"], "approval does not cover appraisal")
    for k in ("actor", "recorded_at", "message"):
        text(approval.get(k), "approval." + k)
    receipt(root, approval["evidence"])
    # Receipt is provenance, not a cryptographic proof of human identity.
    baseline = root / "provenance/design" / sha(proposal_path)
    for name, content in p["documents"].items():
        require(not (root / name).exists() or name in ROOT_DOCS[:2], "would overwrite formal document: " + name)
    for name, content in p["documents"].items():
        atomic(baseline / name, content)
        atomic(root / name, content)
    atomic(root / "registry/tasks.tsv", tasks_table(p["tasks"]))
    canonical = [{k: t[k] for k in TASK_FIELDS} for t in p["tasks"]]
    atomic(root / "registry/plan_tasks.json", canonical)
    # Use this installed source explicitly, without generating a second implementation.
    scripts = str(Path(__file__).resolve().parent)
    hook = "import sys\nsys.path.insert(0, " + repr(scripts) + ")\nfrom project_events import main\nmain()\n"
    atomic(root / "scripts/hooks/project_event.py", hook)
    design = dict(policy="plan_results_v1", version=p["version"],
                  proposal=descriptor(proposal_path, root), approval=descriptor(approval_path, root),
                  appraisal=p["appraisal"], approval_evidence=approval["evidence"],
                  intake=descriptor(root / "provenance/intake.json", root),
                  documents={name: descriptor(baseline / name, root) for name in ROOT_DOCS},
                  active_scope={name: sha(root / name) for name in ROOT_DOCS[:2]},
                  task_manifest=descriptor(root / "registry/plan_tasks.json", root))
    atomic(root / "registry/design.json", design)
    from project_events import emit
    for kind, ref in (('research.appraisal.completed', design['appraisal']),
                      ('research.appraisal.discussed', design['approval'])):
        emit(root, dict(event_type=kind, **ref, verified=True, message='Evidence recorded for '+p['version'],
                        meta=dict(plan_version=p['version'])))
    return verify_design(root)


def verify_design(root):
    root = Path(root).resolve()
    d = read(root / "registry/design.json")
    require(d.get("policy") == "plan_results_v1", "unsupported workflow policy")
    for k in ("proposal", "approval", "appraisal", "approval_evidence", "intake", "task_manifest"):
        receipt(root, d[k])
    proposal = read(receipt(root, d["proposal"]))
    approval = read(receipt(root, d["approval"]))
    require(approval.get("decision") == "approved" and approval.get("proposal_sha256") == d["proposal"]["sha256"]
            and approval.get("appraisal_sha256") == d["appraisal"]["sha256"], "approval linkage changed")
    require(proposal["appraisal"] == d["appraisal"] and proposal["version"] == d["version"],
            "design/proposal mismatch")
    validate_appraisal(root, d["appraisal"])
    for name in ROOT_DOCS:
        require((root / name).is_file(), "missing root document: " + name)
        baseline = receipt(root, d["documents"][name])
        require(baseline.read_text() == proposal["documents"][name], "baseline differs from approved proposal")
    for name in ROOT_DOCS[:2]:
        require(sha(root / name) == d["active_scope"][name] == d["documents"][name]["sha256"],
                "frozen scientific document changed: " + name)
    expected = read(receipt(root, d["task_manifest"]))
    require(expected == [{k: t[k] for k in TASK_FIELDS} for t in proposal["tasks"]], "task manifest changed")
    rows = registry(root)
    require([{k: r[k] for k in TASK_FIELDS} for r in rows] == expected, "plan/registry mapping changed")
    validate_tasks(root, rows, (root / ROOT_DOCS[0]).read_text(), (root / ROOT_DOCS[1]).read_text())
    return dict(policy=d["policy"], version=d["version"], tasks=len(rows), passed=True)


def task(root, task_id):
    rows = registry(root)
    matches = [r for r in rows if r["task_id"] == task_id]
    require(len(matches) == 1, "unknown task: " + task_id)
    return matches[0]


def report_content(root, task_id, manifest_path, run_id=None):
    root = Path(root).resolve()
    design = verify_design(root)
    t = task(root, task_id)
    result = inside(root, t["result_path"])
    manifest_path = Path(manifest_path).resolve()
    require(manifest_path.is_relative_to(result), "report manifest outside task")
    m = read(manifest_path)
    require(m.get("task_id") == task_id and m.get("plan_version") == design["version"], "report task/plan mismatch")
    text(m.get("run_id"), "report.run_id")
    if run_id is not None:
        require(m["run_id"] == run_id, "report run_id mismatch")
    require(m.get("language") in HEADINGS, "report language must be zh or en")
    for k in ("title", "host", "environment"):
        text(m.get(k), k)
    require(isinstance(m.get("command"), list) and m["command"] and
            all(isinstance(x, str) for x in m["command"]), "actual command argv required")
    require(m.get("software") and isinstance(m["software"], list), "software records required")
    for s in m["software"]:
        for k in ("name", "version", "usage"):
            text(s.get(k), "software." + k)
    for k in SECTIONS:
        text(m.get("sections", {}).get(k), "section." + k)
    require(m.get("inputs"), "hashed input provenance required")
    for inp in m["inputs"]:
        for k in ("source", "version", "access"):
            text(inp.get(k), "input." + k)
        receipt(root, inp)
    require(m.get("tables") and m.get("figures"), "nonempty table and figure lists required")
    tables = {}
    for table in m["tables"]:
        p = receipt(result, table)
        text(table.get("caption"), "table caption")
        with p.open(encoding="utf-8", newline="") as f:
            rows = list(csv.reader(f, delimiter="\t" if p.suffix == ".tsv" else ","))
        require(len(rows) > 1 and rows[0] and all(len(r) == len(rows[0]) for r in rows), "invalid/empty table")
        require(len(rows) <= 1001, "report table must be a compact summary (at most 1000 data rows)")
        require(table["path"] not in tables, "duplicate report table")
        tables[table["path"]] = rows
    for fig in m["figures"]:
        p = receipt(result, fig)
        text(fig.get("caption"), "figure caption")
        text(fig.get("method"), "figure generation method")
        require(p.suffix.lower() in (".png", ".jpg", ".jpeg"), "report figure must be PNG or JPEG")
        require(fig.get("source_tables") and set(fig["source_tables"]) <= set(tables), "figure/table lineage missing")
    return t, result, m, tables


def render_report(root, task_id, manifest_path):
    from docx import Document
    from docx.shared import Inches, Mm, Pt
    from docx.oxml.ns import qn
    t, result, m, tables = report_content(root, task_id, manifest_path)
    require(not (result / "_evidence/report_receipt.json").exists(),
            "report already rendered; preserve prior run and use reviewed report revision")
    d = Document()
    # python-docx's bundled template has a zoom element without the required
    # percent attribute. Normalize it so generated OOXML validates strictly.
    from docx.oxml import OxmlElement
    zoom = d.settings.element.find(qn('w:zoom'))
    if zoom is None:
        zoom = OxmlElement('w:zoom')
        d.settings.element.append(zoom)
    zoom.set(qn('w:percent'), '100')
    section = d.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    for name in ('Normal', 'Title', 'Heading 1'):
        style = d.styles[name]
        style.font.name = 'Arial'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    d.styles['Normal'].font.size = Pt(10)
    d.add_heading(m["title"], 0)
    md = ["# " + m["title"], ""]
    keys = ('task_id', 'plan_version', 'run_id', 'host', 'environment', 'command')
    labels = ('任务编号', '计划版本', '运行编号', '执行主机', '环境', '实际命令') if m['language'] == 'zh' else keys
    import shlex
    for key, label in zip(keys, labels):
        value = shlex.join(m[key]) if key == 'command' else m[key]
        line = label + ': ' + value
        d.add_paragraph(line)
        md += [line, '']
    root = Path(root).resolve()
    links = [(name, os.path.relpath(root / name, result)) for name in ROOT_DOCS[:2]]
    md += [" | ".join("[" + name + "](" + rel + ")" for name, rel in links), ""]
    for name, rel in links:
        d.add_paragraph(name + ": " + rel)
    # Both formats are rendered from exactly the same supplied content.
    for key, heading in zip(SECTIONS, HEADINGS[m["language"]]):
        d.add_heading(heading, 1)
        value = m["sections"][key]
        for line in value.splitlines():
            d.add_paragraph(line)
        md += ["## " + heading, "", value, ""]
    for record in m["inputs"] + m["software"]:
        for key, value in record.items():
            label = {'path': '路径', 'sha256': 'SHA-256', 'source': '来源', 'version': '版本',
                     'access': '获取方式', 'name': '名称', 'usage': '用途'}.get(key, key) if m['language'] == 'zh' else key
            line = label + ': ' + str(value)
            d.add_paragraph(line)
            md += [line, '']
    for item in m["tables"]:
        d.add_paragraph(item["caption"])
        md += [item["caption"], ""]
        rows = tables[item["path"]]
        word_table = d.add_table(rows=0, cols=len(rows[0]))
        word_table.style = 'Table Grid'
        word_table.autofit = False
        for column in word_table.columns:
            column.width = Mm(166 / len(rows[0]))
        for row in rows:
            cells = word_table.add_row().cells
            for c, value in zip(cells, row):
                c.width = Mm(166 / len(rows[0]))
                c.text = value
        md += [table_line(rows[0]), table_line(["---"] * len(rows[0]))] + [table_line(r) for r in rows[1:]] + [""]
    for item in m["figures"]:
        d.add_picture(str(result / item["path"]), width=Inches(5.5))
        d.add_paragraph(item["caption"])
        md += ["![" + item["caption"] + "](" + item["path"] + ")", ""]
    atomic(result / "README.md", "\n".join(md) + "\n")
    fd, temp = tempfile.mkstemp(suffix=".docx", dir=result)
    os.close(fd)
    try:
        d.save(temp)
        os.replace(temp, result / "REPORT.docx")
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
    outputs = ["README.md", "REPORT.docx"] + [x["path"] for x in m["tables"] + m["figures"]]
    r = dict(task_id=task_id, run_id=m["run_id"], plan_version=m["plan_version"],
             manifest=descriptor(manifest_path, result),
             artifacts=[descriptor(result / p, result) for p in outputs])
    atomic(result / "_evidence/report_receipt.json", r)
    return r


def validate_package(root, task_id, result_path, run_id):
    from docx import Document
    t = task(root, task_id)
    result = inside(root, result_path)
    require(result == inside(root, t["result_path"]), "formal report must use canonical task path")
    r = read(result / "_evidence/report_receipt.json")
    require(r["task_id"] == task_id and r["run_id"] == run_id, "report receipt run/task mismatch")
    mp = receipt(result, r["manifest"])
    _, _, m, tables = report_content(root, task_id, mp, run_id)
    for a in r["artifacts"]:
        receipt(result, a)
    paths = {a["path"] for a in r["artifacts"]}
    required = {"README.md", "REPORT.docx"} | {x["path"] for x in m["tables"] + m["figures"]}
    require(required <= paths, "unhashed report output")
    markdown = (result / "README.md").read_text()
    word = Document(str(result / "REPORT.docx"))
    paragraphs = "\n".join(p.text for p in word.paragraphs)
    for value in m["sections"].values():
        require(value in markdown and value in paragraphs, "Markdown/Word content mismatch")
    require(run_id in markdown and run_id in paragraphs, "run identity absent in report")
    expected = list(tables.values())
    actual = [[[c.text for c in row.cells] for row in tab.rows] for tab in word.tables]
    require(actual == expected, "Word/source table mismatch")
    for rows in expected:
        for row in rows:
            rendered = table_line(row)
            require(rendered in markdown, 'Markdown/source table mismatch')
    for figure in m['figures']:
        require('![' + figure['caption'] + '](' + figure['path'] + ')' in markdown,
                'Markdown/source figure mismatch')
    with zipfile.ZipFile(result / "REPORT.docx") as z:
        embedded = {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith("word/media/")}
    require({x["sha256"] for x in m["figures"]} <= embedded, "Word/source figure mismatch")
    # QC must cover every final deliverable, the manifest, receipt and reflection.
    qc = read(result / "_evidence/qc.json")
    require(qc.get("passed") is True and qc.get("run_id") == run_id, "QC run mismatch")
    need = required | {r["manifest"]["path"], "_evidence/report_receipt.json", "_evidence/reflection.json"}
    actual_qc = {a["path"] for a in qc.get("artifacts", [])}
    require(need <= actual_qc, "QC does not cover complete report evidence")
    for a in qc["artifacts"]:
        receipt(result, a)
    return dict(passed=True, task_id=task_id, run_id=run_id, artifacts=len(need))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("mode", choices=("draft", "freeze", "check", "render", "validate-package"))
    ap.add_argument("--root", type=Path, required=True)
    ap.add_argument("--spec", type=Path)
    ap.add_argument("--proposal", type=Path)
    ap.add_argument("--approval", type=Path)
    ap.add_argument("--task-id")
    ap.add_argument("--manifest", type=Path)
    ap.add_argument("--run-id")
    a = ap.parse_args()
    if a.mode == "draft":
        out = draft(a.root, read(a.spec))
    elif a.mode == "freeze":
        out = freeze(a.root, a.proposal, a.approval)
    elif a.mode == "check":
        out = verify_design(a.root)
    elif a.mode == "render":
        out = render_report(a.root, a.task_id, a.manifest)
    else:
        out = validate_package(a.root, a.task_id, task(a.root, a.task_id)["result_path"], a.run_id)
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
